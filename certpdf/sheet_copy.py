"""Helpers for copying tabs out of the original workbooks.

Values are the ones Excel last calculated and saved, so the copies need
neither the rest of the original workbook nor any recalculation.
"""

from __future__ import annotations

import re
import warnings
import zipfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles.fills import PatternFill
from openpyxl.utils.indexed_list import IndexedList

_PAGE_SETUP_ATTRS = ("orientation", "paperSize", "scale", "fitToWidth", "fitToHeight", "firstPageNumber",
                     "useFirstPageNumber", "pageOrder", "blackAndWhite", "draft", "cellComments", "errors")


def use_default_font(target_book, font) -> None:
    """Give a new workbook the source's default font.

    Excel measures column widths in characters of the default font, so with a
    different one (a new workbook has Calibri 11, the templates Arial 10) every
    column changes width and pages break differently. Call before adding cells.
    """
    if len(target_book._fonts) != 1:
        return  # cells already use other fonts; changing the list would re-point them
    target_book._fonts = IndexedList([copy(font)])
    target_book._named_styles["Normal"].font = copy(font)


def _without_empty_fill(dxf):
    """Drop a conditional format's "no fill" setting.

    openpyxl reads patternType="none" as "not set" and writes it back without
    it, and Excel then shows a *solid* fill in the automatic colour: black
    boxes. A fill without a pattern and without a real colour means "no fill".
    """
    fill = getattr(dxf, "fill", None)
    if isinstance(fill, PatternFill) and fill.patternType is None:
        colour = fill.bgColor
        if colour is None or colour.type == "auto" or (colour.type == "rgb" and colour.rgb == "00000000") \
                or (colour.type == "indexed" and colour.indexed in (64, 65)):
            dxf.fill = None
    return dxf


_SPACES_ONLY = re.compile(rb"<t>(\s+)</t>")


def keep_spaces(path: Path) -> None:
    """Mark text that is only spaces as kept, in a saved workbook.

    openpyxl writes such text (a rich-text piece like the " " in "Client" + " "
    + "Name:", or a cell holding a space) without xml:space="preserve", and
    Excel then drops the spaces: "ClientName:", or a cell that no longer stops
    the text beside it from running over.
    """
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info)) for info in source.infolist()]
    parts = [(info, _SPACES_ONLY.sub(rb'<t xml:space="preserve">\1</t>', data))
             if info.filename.startswith("xl/worksheets/") or info.filename == "xl/sharedStrings.xml"
             else (info, data) for info, data in entries]
    if parts == entries:
        return
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for info, data in parts:
            target.writestr(info, data)


def load_values(path: Path):
    """Open a workbook with the values Excel saved (not the formulas)."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl warns about features it doesn't need to read
        return load_workbook(path, data_only=True, rich_text=True, keep_links=False)
