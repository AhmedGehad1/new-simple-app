"""Put many copies of a tab (one per device) one after another on a single sheet.

Every tab ("block") keeps its own column widths: the columns of all blocks
are laid over each other and the sheet gets a finer grid with a column
boundary wherever any block has one. A block's column then covers one or
more grid columns, and cells that need their original width are merged
across them. Widths are worked out in screen pixels, so they come out exact.

One sheet has one print scale, so each block is enlarged by (its own scale /
the common scale) - column widths, row heights and font sizes - and prints
at the size it had. A page break follows every block.
"""

from __future__ import annotations

import re
from bisect import bisect_left
from copy import copy
from dataclasses import dataclass, field

from openpyxl.cell.cell import MergedCell
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.styles import Alignment, Font, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.pagebreak import Break

from .sheet_copy import _PAGE_SETUP_ATTRS, _without_empty_fill

DIGIT_PX = 7          # width of a digit in pixels for the templates' default font (Arial 10 / Calibri 11)
POINTS_PER_PX = 0.75
PAPER_POINTS = {1: (612.0, 792.0), 9: (595.3, 841.9)}  # Letter, A4
A4 = 9
FOOTER_HEIGHT = 12.75  # points, for a footer printed as a line of text
_CELL_REF = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
_NO_SIDE = Side()


def width_px(chars: float) -> int:
    """Excel column width (in characters) to pixels."""
    return int(chars * DIGIT_PX + 0.07) if chars > 0 else 0


@dataclass
class Block:
    """What is needed of one tab's printed area, detached from its workbook."""
    label: str
    widths_px: list[int]                     # per column of the printed area
    heights: list[float]                     # per row, points
    hidden_rows: set[int]
    styles: list[tuple]                      # (font, fill, border, number_format, protection, alignment)
    cells: list[tuple]                       # (row, col, value, style index or None), 0-based in the block
    merges: list[tuple]                      # (row1, col1, row2, col2), 0-based
    cut_sides: dict                          # (row, col) -> sides where the print area cuts a merged box
    formats: list[tuple]                     # ([(row1, col1, row2, col2)...], rule)
    breaks: list[int]                        # page break after these rows (0-based)
    origin: tuple[int, int]                  # (first row, first column) of the printed area in the tab
    scale: int                               # print scale the tab prints at
    margins: object
    orientation: str | None
    paper: int | None
    page_setup: dict                         # the tab's other page setup settings
    footer: object
    print_options: object
    show_grid: bool | None
    view: str | None
    scaled: float = 1.0                      # set when laying out: this block's scale / the common scale
    columns: list[tuple] = field(default_factory=list)  # grid columns (first, last) for each block column
    first_row: int = 0                       # sheet row of the block's first row


# --------------------------------------------------------------------------
# Reading a tab
# --------------------------------------------------------------------------

def printed_area(ws) -> tuple[int, int, int, int]:
    """(first column, first row, last column, last row) of what the tab prints."""
    if ws.print_area:
        boxes = [range_boundaries(ref.replace("$", "")) for ref in re.findall(r"!([$A-Za-z0-9:]+)", ws.print_area)]
        boxes = [b for b in boxes if None not in b]
        if boxes:
            return (min(b[0] for b in boxes), min(b[1] for b in boxes),
                    max(b[2] for b in boxes), max(b[3] for b in boxes))
    # No print area: Excel prints from A1 to the last cell that shows something - a value, a border or
    # a fill. Cells that only have a font set don't count; on a centred page they would move everything.
    last_col = last_row = 0
    for cell in list(ws._cells.values()):
        if cell.value is not None or _shows(cell):
            last_col, last_row = max(last_col, cell.column), max(last_row, cell.row)
    for merged in ws.merged_cells.ranges:
        anchor = ws._cells.get((merged.min_row, merged.min_col))
        if anchor is not None and (anchor.value is not None or _shows(anchor)):
            last_col, last_row = max(last_col, merged.max_col), max(last_row, merged.max_row)
    if not last_col:
        return range_boundaries(ws.calculate_dimension())
    return 1, 1, last_col, last_row


def _shows(cell) -> bool:
    """True when an empty cell still prints something: a border or a fill."""
    if not cell.has_style:
        return False
    border = cell.border
    if any(side is not None and side.style for side in (border.left, border.right, border.top, border.bottom)):
        return True
    if border.diagonal is not None and border.diagonal.style and (border.diagonalUp or border.diagonalDown):
        return True
    return getattr(cell.fill, "fill_type", "gradient") not in (None, "none")  # a gradient fill has no type


def _column_widths(ws, first: int, last: int) -> list[float]:
    fmt = ws.sheet_format
    default = fmt.defaultColWidth if fmt.defaultColWidth else (fmt.baseColWidth or 8) + 0.43
    widths = [default] * (last - first + 1)
    for dim in list(ws.column_dimensions.values()):
        low, high = dim.min or 0, dim.max or 0
        for col in range(max(low, first), min(high, last) + 1):
            widths[col - first] = 0.0 if dim.hidden else (dim.width if dim.width is not None else default)
    return widths


def _effective_scale(ws, width_pt: float, height_pt: float) -> int:
    setup = ws.page_setup
    scale = int(setup.scale or 100)
    fit = ws.sheet_properties.pageSetUpPr
    if fit is not None and fit.fitToPage and setup.scale and setup.scale < 100:
        # With "Fit to page" Excel saves the scale it worked out, so that is what the tab prints at.
        return int(setup.scale)
    if fit is not None and fit.fitToPage:
        paper_w, paper_h = PAPER_POINTS.get(setup.paperSize or A4, PAPER_POINTS[A4])
        if setup.orientation == "landscape":
            paper_w, paper_h = paper_h, paper_w
        m = ws.page_margins
        room_w = paper_w - 72 * ((m.left or 0) + (m.right or 0))
        room_h = paper_h - 72 * ((m.top or 0) + (m.bottom or 0))
        wide = 1 if setup.fitToWidth is None else setup.fitToWidth
        tall = 1 if setup.fitToHeight is None else setup.fitToHeight
        fitted = 100.0
        if wide and width_pt:
            fitted = min(fitted, 100 * wide * room_w / width_pt)
        if tall and height_pt:
            fitted = min(fitted, 100 * tall * room_h / height_pt)
        scale = max(10, int(fitted))
    return scale


def read_block(ws, label: str) -> Block:
    first_col, first_row, last_col, last_row = printed_area(ws)
    widths = [width_px(w) for w in _column_widths(ws, first_col, last_col)]
    default_height = ws.sheet_format.defaultRowHeight or 15
    heights, hidden = [], set()
    for row in range(first_row, last_row + 1):
        dim = ws.row_dimensions.get(row)
        heights.append(dim.height if dim is not None and dim.height is not None else default_height)
        if dim is not None and dim.hidden:
            hidden.add(row - first_row)

    styles: list[tuple] = []
    style_index: dict[tuple, int] = {}
    cells = []
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, min_col=first_col, max_col=last_col):
        for cell in row:
            value = getattr(cell, "value", None)
            if value is None and not cell.has_style:
                continue
            ref = None
            if cell.has_style:
                key = tuple(cell._style)
                if key not in style_index:
                    style_index[key] = len(styles)
                    styles.append((copy(cell.font), copy(cell.fill), copy(cell.border), cell.number_format,
                                   copy(cell.protection), copy(cell.alignment)))
                ref = style_index[key]
            cells.append((cell.row - first_row, cell.column - first_col, value, ref))

    merges = []
    cut_sides: dict[tuple, frozenset] = {}
    for merged in ws.merged_cells.ranges:
        r1, c1 = max(merged.min_row, first_row), max(merged.min_col, first_col)
        r2, c2 = min(merged.max_row, last_row), min(merged.max_col, last_col)
        if r1 <= r2 and c1 <= c2:
            merges.append((r1 - first_row, c1 - first_col, r2 - first_row, c2 - first_col))
            # Where the print area cuts through a merged box, the edge is inside the box in the
            # original, so no border is printed there.
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    sides = set()
                    if r == r2 and merged.max_row > r2:
                        sides.add("bottom")
                    if r == r1 and merged.min_row < r1:
                        sides.add("top")
                    if c == c2 and merged.max_col > c2:
                        sides.add("right")
                    if c == c1 and merged.min_col < c1:
                        sides.add("left")
                    if sides:
                        cut_sides[(r - first_row, c - first_col)] = frozenset(sides)

    formats = []
    for formatting in ws.conditional_formatting:
        ranges = []
        for cr in formatting.sqref.ranges:
            r1, c1 = max(cr.min_row, first_row), max(cr.min_col, first_col)
            r2, c2 = min(cr.max_row, last_row), min(cr.max_col, last_col)
            if r1 <= r2 and c1 <= c2:
                ranges.append((r1 - first_row, c1 - first_col, r2 - first_row, c2 - first_col))
        if ranges:
            for rule in formatting.rules:
                formats.append((ranges, rule))

    breaks = sorted({brk.id - first_row for brk in ws.row_breaks.brk if first_row <= brk.id < last_row})
    width_pt = sum(widths) * POINTS_PER_PX
    height_pt = sum(h for i, h in enumerate(heights) if i not in hidden)
    return Block(
        label=label, widths_px=widths, heights=heights, hidden_rows=hidden, styles=styles, cells=cells,
        merges=merges, cut_sides=cut_sides, formats=formats, breaks=breaks, origin=(first_row, first_col),
        scale=_effective_scale(ws, width_pt, height_pt), margins=copy(ws.page_margins),
        orientation=ws.page_setup.orientation, paper=ws.page_setup.paperSize,
        page_setup={attr: getattr(ws.page_setup, attr) for attr in _PAGE_SETUP_ATTRS},
        footer=copy(ws.HeaderFooter),
        print_options=copy(ws.print_options), show_grid=ws.sheet_view.showGridLines, view=ws.sheet_view.view)


# --------------------------------------------------------------------------
# Writing the combined sheet
# --------------------------------------------------------------------------

class _Styles:
    """Turns a block's cell style into a style of the target workbook, once per variant."""

    def __init__(self) -> None:
        self._cache: dict[tuple, object] = {}

    def apply(self, cell, block_id: int, block: Block, ref: int, first: bool, last: bool,
              cut: frozenset = frozenset()) -> None:
        key = (block_id, ref, first, last, cut)
        cached = self._cache.get(key)
        if cached is not None:
            cell._style = copy(cached)
            return
        font, fill, border, number_format, protection, alignment = block.styles[ref]
        if abs(block.scaled - 1) > 0.02 and font.sz:
            font = copy(font)
            font.sz = round(font.sz * block.scaled * 2) / 2
        if not (first and last) or cut:
            # A column split over several grid columns: only its outer edges have left/right borders.
            border = copy(border)
            if not first or "left" in cut:
                border.left = _NO_SIDE
            if not last or "right" in cut:
                border.right = _NO_SIDE
            if "top" in cut:
                border.top = _NO_SIDE
            if "bottom" in cut:
                border.bottom = _NO_SIDE
        cell.font, cell.fill, cell.border = copy(font), copy(fill), border
        cell.number_format, cell.protection, cell.alignment = number_format, copy(protection), copy(alignment)
        self._cache[key] = copy(cell._style)


def _needs_own_width(value, alignment) -> bool:
    """Whether text must stay inside its original column width (so gets merged across grid columns).

    Left-aligned text is left unmerged so it can still run into empty neighbouring cells.
    """
    if value is None:
        return False
    horizontal = getattr(alignment, "horizontal", None)
    if horizontal in ("center", "centerContinuous", "right", "justify", "distributed", "fill"):
        return True
    if getattr(alignment, "wrap_text", False) or getattr(alignment, "shrink_to_fit", False):
        return True
    return horizontal in (None, "general") and not isinstance(value, str) and not hasattr(value, "__iter__")


def _translate(formula: str, block: Block, row_of) -> str | None:
    """Move the cell references of a conditional format formula to where the block now is."""
    tokens = Tokenizer("=" + formula)
    for token in tokens.items:
        if token.type != Token.OPERAND or token.subtype != Token.RANGE:
            continue
        if "!" in token.value:
            return None  # refers to another tab, which isn't in the combined file
        parts = []
        for number, part in enumerate(token.value.split(":")):
            match = _CELL_REF.match(part)
            if not match:
                return None
            col_abs, letters, row_abs, digits = match.groups()
            col = _column_index(letters) - block.origin[1]
            row = int(digits) - block.origin[0]
            if not (0 <= col < len(block.columns)):
                return None
            grid_col = block.columns[col][1 if number else 0]
            parts.append(f"{col_abs}{get_column_letter(grid_col)}{row_abs}{row_of(row)}")
        token.value = ":".join(parts)
    return tokens.render()[1:]


def _column_index(letters: str) -> int:
    number = 0
    for char in letters.upper():
        number = number * 26 + ord(char) - 64
    return number


def write_blocks(ws, blocks: list[Block]) -> list[str]:
    """Lay *blocks* out one after another on *ws*. Returns notes for the user."""
    notes = []
    common = min(block.scale for block in blocks)
    for block in blocks:
        # Scales within 2 % of each other (Excel's fit-to-page gives 87 % for one device and 88 %
        # for the next) are left alone: resizing would split columns for a difference nobody sees.
        block.scaled = block.scale / common if block.scale / common > 1.02 else 1.0

    # Shared column grid (pixels). Narrower blocks are centred on the widest when the tab
    # prints centred, otherwise they start at the left like the original.
    centred = bool(blocks[0].print_options.horizontalCentered)
    scaled_widths = [[round(w * b.scaled) for w in b.widths_px] for b in blocks]
    total = max(sum(widths) for widths in scaled_widths)
    edges = {0, total}
    starts = []
    for widths in scaled_widths:
        x = (total - sum(widths)) // 2 if centred else 0
        starts.append(x)
        edges.add(x)
        for w in widths:
            x += w
            edges.add(x)
    grid = sorted(edges)
    position = {x: i for i, x in enumerate(grid)}
    for i in range(len(grid) - 1):
        ws.column_dimensions[get_column_letter(i + 1)].width = (grid[i + 1] - grid[i]) / DIGIT_PX
    for block, widths, x in zip(blocks, scaled_widths, starts):
        block.columns = []
        for w in widths:
            block.columns.append((position[x] + 1, position[x + w]))  # empty when first > last (hidden)
            x += w

    first = blocks[0]
    margins = ws.page_margins
    for side in ("left", "right", "top", "bottom", "header", "footer"):
        values = [getattr(b.margins, side) for b in blocks if getattr(b.margins, side) is not None]
        if values:
            setattr(margins, side, min(values))
    paper_w, paper_h = PAPER_POINTS.get(first.paper or A4, PAPER_POINTS[A4])
    if first.orientation == "landscape":
        paper_w, paper_h = paper_h, paper_w
    page_room = paper_h - 72 * ((margins.top or 0) + (margins.bottom or 0))

    # One sheet has one page footer. When the tabs' footers differ, each tab's
    # footer becomes a line of text at the bottom of its pages (where it fits).
    shared_footer = len({_footer_text(b.footer) for b in blocks}) == 1
    footer_lines = 0

    styles = _Styles()
    next_row = 1
    priority = 0
    for block_id, block in enumerate(blocks):
        block.first_row = next_row
        line = "" if shared_footer else _footer_line(block.footer)
        after = []  # block rows after which a footer line is inserted
        if line:
            ends = [*block.breaks, len(block.heights) - 1]
            start = 0
            for end in ends:
                printed = sum(h for r, h in enumerate(block.heights[start:end + 1], start)
                              if r not in block.hidden_rows) * block.scale / 100
                if printed + FOOTER_HEIGHT * block.scale / 100 <= page_room:
                    after.append(end)
                start = end + 1
            footer_lines += len(after)

        def row_of(r, _start=next_row, _after=tuple(after)):
            return _start + r + bisect_left(_after, r)

        for r, height in enumerate(block.heights):
            dim = ws.row_dimensions[row_of(r)]
            dim.height = round(height * block.scaled * 4) / 4
            if r in block.hidden_rows:
                dim.hidden = True
        span = (block.columns[0][0], block.columns[-1][1])
        for r in after:
            footer_row = row_of(r) + 1
            ws.row_dimensions[footer_row].height = round(FOOTER_HEIGHT * block.scaled * 4) / 4
            if span[0] < span[1]:
                ws.merge_cells(start_row=footer_row, start_column=span[0], end_row=footer_row, end_column=span[1])
            cell = ws.cell(row=footer_row, column=span[0], value=line)
            cell.font = Font(size=round(9 * block.scaled * 2) / 2)
            cell.alignment = Alignment(horizontal="center", vertical="bottom")

        # Merge first: merging resets the covered cells, whose styles are set below.
        merged_anchor = {}
        for r1, c1, r2, c2 in block.merges:
            g1, g2 = block.columns[c1][0], block.columns[c2][1]
            if g1 <= g2 and (r1 != r2 or g1 != g2):
                ws.merge_cells(start_row=row_of(r1), start_column=g1, end_row=row_of(r2), end_column=g2)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    merged_anchor[(r, c)] = (r, c) == (r1, c1)
        for r, c, value, ref in block.cells:
            g1, g2 = block.columns[c]
            if g1 > g2:
                continue  # hidden column
            alignment = block.styles[ref][5] if ref is not None else None
            if (r, c) not in merged_anchor and g1 < g2 and _needs_own_width(value, alignment):
                ws.merge_cells(start_row=row_of(r), start_column=g1, end_row=row_of(r), end_column=g2)
        for r, c, value, ref in block.cells:
            g1, g2 = block.columns[c]
            for g in range(g1, g2 + 1):
                cell = ws.cell(row=row_of(r), column=g)
                if ref is not None:
                    styles.apply(cell, block_id, block, ref, g == g1, g == g2, block.cut_sides.get((r, c), frozenset()))
                if g == g1 and value is not None and not isinstance(cell, MergedCell):
                    cell.value = value
                    if isinstance(value, str) and value.startswith("="):
                        cell.data_type = "s"  # text that only looks like a formula

        for ranges, rule in sorted(block.formats, key=lambda item: item[1].priority or 0):
            new_rule = copy(rule)
            new_rule.dxf = _without_empty_fill(copy(rule.dxf))
            priority += 1
            new_rule.priority = priority
            if rule.formula:
                formulas = [_translate(f, block, row_of) for f in rule.formula]
                if any(f is None for f in formulas):
                    continue
                new_rule.formula = formulas
            refs = []
            for r1, c1, r2, c2 in ranges:
                g1, g2 = block.columns[c1][0], block.columns[c2][1]
                if g1 <= g2:
                    refs.append(f"{get_column_letter(g1)}{row_of(r1)}:{get_column_letter(g2)}{row_of(r2)}")
            if refs:
                ws.conditional_formatting.add(" ".join(refs), new_rule)

        for r in block.breaks:
            ws.row_breaks.append(Break(id=row_of(r) + (1 if r in after else 0)))
        last = len(block.heights) - 1
        next_row = row_of(last) + (2 if last in after else 1)
        if block_id < len(blocks) - 1:
            ws.row_breaks.append(Break(id=next_row - 1))

    # Page setup for the whole sheet
    setup = ws.page_setup
    # Exactly the tab's own page setup (paper, orientation, ...), except that one scale replaces
    # "fit to page": Excel ignores page breaks when fitting.
    for attr, value in first.page_setup.items():
        setattr(setup, attr, value)
    setup.scale = common
    setup.fitToWidth = setup.fitToHeight = None
    ws.print_options = copy(first.print_options)
    ws.print_area = f"A1:{get_column_letter(len(grid) - 1)}{next_row - 1}"
    if shared_footer:
        ws.HeaderFooter = copy(first.footer)
    elif footer_lines:
        notes.append("the tabs have different page footers, so each footer is printed as a line "
                     "at the bottom of its own pages")
    ws.sheet_view.showGridLines = first.show_grid
    ws.sheet_view.view = first.view
    if any(b.scaled != 1 for b in blocks):
        notes.append(f"printed at {common}%; tabs that printed larger were enlarged to match")
    return notes


def _footer_line(header_footer) -> str:
    """The odd-page footer as one line of text, without Excel's &-codes (page numbers, fonts...)."""
    item = getattr(header_footer, "oddFooter", None)
    if item is None:
        return ""
    parts = []
    for section in ("left", "center", "right"):
        text = getattr(getattr(item, section, None), "text", None) or ""
        text = re.sub(r'&"[^"]*"|&\d+|&[A-Za-z&]', "", text).strip()
        if text:
            parts.append(text)
    return "      ".join(parts)


def _footer_text(header_footer) -> str:
    parts = []
    for name in ("oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter"):
        item = getattr(header_footer, name, None)
        for section in ("left", "center", "right"):
            text = getattr(getattr(item, section, None), "text", None) if item is not None else None
            parts.append(text or "")
    return "|".join(parts).strip("|")
