"""Combine the wanted tabs of every workbook into Excel files, one tab per sheet.

Works on the files themselves (openpyxl): it needs neither Excel nor
LibreOffice, and never prints. Each copied tab keeps its layout and page
setup, and formulas are replaced by the values Excel last calculated, so the
copies don't depend on the other tabs of the original workbook.
"""

from __future__ import annotations

import os
import threading
from pathlib import Path

from openpyxl import Workbook

from .builder import BuildResult, Cancelled, OutputError, OutputResult, Reporter, WorkbookResult
from .config import GROUP_TITLES, GROUPS, SHEET_SPECS
from .engines import EngineError, LibreOfficeEngine
from .sheet_copy import adopt_default_font, copy_sheet, load_values, tab_title
from .workbook_info import match_sheets

ENGINE_NAME = "Excel files"
TAB_LABELS = {
    "ecg_cert": "ECG cert",
    "nibp_cert": "NIBP cert",
    "ecg_pf": "ECG pass-fail",
    "nibp_pf": "NIBP pass-fail",
}


class _Sources:
    """Opens source workbooks; old .xls files are converted with LibreOffice when it's installed."""

    def __init__(self) -> None:
        self._libreoffice: LibreOfficeEngine | None = None

    def load(self, path: Path):
        if path.suffix.lower() == ".xls":
            path = self._convert_xls(path)
        try:
            return load_values(path)
        except Exception as exc:
            raise EngineError(f"the file could not be read ({exc or type(exc).__name__})") from None

    def _convert_xls(self, path: Path) -> Path:
        if self._libreoffice is None:
            if not LibreOfficeEngine.is_available():
                raise EngineError("Excel files can only be made from .xlsx files; this is an old .xls file "
                                  "(install LibreOffice to convert it automatically)")
            self._libreoffice = LibreOfficeEngine()
            self._libreoffice.start()
        return self._libreoffice.convert_to_xlsx(path)

    def close(self) -> None:
        if self._libreoffice is not None:
            self._libreoffice.stop()
            self._libreoffice = None


def build_workbooks(
    files: list[Path],
    sheet_names: dict[str, str],
    outputs: dict[str, Path],
    reporter: Reporter,
    cancel: threading.Event | None = None,
) -> BuildResult:
    """Copy the tabs of *files* into one Excel file per group in *outputs*.

    Tabs are ordered workbook by workbook, and inside each workbook in the
    order given by GROUPS (ECG first, then NIBP).
    """
    cancel = cancel or threading.Event()
    wanted = {spec.key: sheet_names.get(spec.key) or spec.default_name for spec in SHEET_SPECS}
    books = {group: Workbook() for group in outputs}
    for book in books.values():
        book.remove(book.active)
    used_titles: dict[str, set[str]] = {group: set() for group in outputs}
    counts = {group: 0 for group in outputs}
    results: list[WorkbookResult] = []
    total_steps = len(files) + 1
    sources = _Sources()

    reporter.log("info", f"Making Excel files from {len(files)} workbook(s).")
    try:
        for index, path in enumerate(files):
            if cancel.is_set():
                raise Cancelled()
            path = Path(path)
            reporter.progress(index, total_steps, f"Excel files: copying {path.name}  ({index + 1} of {len(files)})")
            reporter.file_status(index, "working", "Copying tabs…")
            result = WorkbookResult(path)
            results.append(result)
            try:
                source = sources.load(path)
            except EngineError as exc:
                result.problems.append(str(exc))
                reporter.log("error", f"{path.name}: {exc}")
                reporter.file_status(index, "error", "Failed – see the messages below")
                continue

            result.found = match_sheets(wanted, source.sheetnames)
            for key, actual in result.found.items():
                if actual is None:
                    reporter.log("warning", f"{path.name}: tab '{wanted[key]}' not found "
                                            f"(tabs in this file: {', '.join(source.sheetnames) or 'none'})")
            copied = []
            for group, keys in GROUPS.items():
                if group not in outputs:
                    continue
                for key in keys:
                    actual = result.found.get(key)
                    if actual is None:
                        continue
                    title = tab_title(path.stem, TAB_LABELS[key], used_titles[group])
                    if counts[group] == 0:
                        adopt_default_font(books[group], source)
                    target = books[group].create_sheet(title)
                    try:
                        copy_sheet(source[actual], target)
                    except Exception as exc:
                        books[group].remove(target)
                        message = f"tab '{actual}' could not be copied: {exc or type(exc).__name__}"
                        result.problems.append(message)
                        reporter.log("error", f"{path.name}: {message}")
                        continue
                    counts[group] += 1
                    copied.append(actual)
            if copied:
                reporter.log("info", f"{path.name}: copied {', '.join(copied)}")
            if result.problems:
                reporter.file_status(index, "error", "Failed – see the messages below")
            elif result.missing:
                reporter.file_status(index, "warning", f"Done – {len(result.missing)} tab(s) missing")
            else:
                reporter.file_status(index, "done", "Done")

        if cancel.is_set():
            raise Cancelled()
        reporter.progress(len(files), total_steps, "Saving Excel files…")
        written: dict[str, OutputResult | None] = {}
        for group, out_path in outputs.items():
            written[group] = _save(group, books[group], counts[group], Path(out_path), reporter)
    finally:
        sources.close()

    reporter.progress(total_steps, total_steps, "Excel files finished")
    return BuildResult(ENGINE_NAME, results, written)


def _save(group: str, book: Workbook, tabs: int, out_path: Path, reporter: Reporter) -> OutputResult | None:
    title = GROUP_TITLES[group]
    if tabs == 0:
        reporter.log("warning", f"{title}: no Excel file saved – none of the files had these tabs.")
        return None
    book.active = 0
    out_path.parent.mkdir(parents=True, exist_ok=True)
    partial = out_path.with_name(out_path.name + ".part")
    try:
        book.save(partial)
        os.replace(partial, out_path)
    except PermissionError:
        _remove_quietly(partial)
        raise OutputError(f"Could not save '{out_path.name}'. If it is open in Excel, "
                          f"close it and try again.") from None
    except OSError as exc:
        _remove_quietly(partial)
        raise OutputError(f"Could not save '{out_path}': {exc.strerror or exc}") from None
    reporter.log("success", f"Saved {out_path.name}: {tabs} tab(s).")
    return OutputResult(out_path, pages=0, tabs=tabs)


def _remove_quietly(path: Path) -> None:
    try:
        os.remove(path)
    except OSError:
        pass
