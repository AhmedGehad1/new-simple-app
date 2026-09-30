from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

openpyxl = pytest.importorskip("openpyxl")

TABS = ["Info", "ECG.certificate", "NIBP.certificate", "ECG.pass-fail", "NIBP.pass-fail", "Raw data"]


def make_workbook(path: Path, device: str, tabs=TABS, hidden=("Raw data",)) -> Path:
    """A workbook whose every tab says which device and tab it is."""
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for tab in tabs:
        ws = wb.create_sheet(tab)
        ws["A1"] = f"{device} {tab}"
        ws["A1"].font = Font(bold=True, size=14)
        for row in range(3, 12):
            ws.cell(row, 1, f"Check {row}")
            ws.cell(row, 2, "PASS")
        if tab in hidden:
            ws.sheet_state = "hidden"
    wb.save(path)
    return path


@pytest.fixture
def workbook_factory(tmp_path):
    def factory(name: str, device: str | None = None, **kwargs) -> Path:
        return make_workbook(tmp_path / name, device or Path(name).stem, **kwargs)
    return factory
