import re
import subprocess
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule, Rule
from openpyxl.styles import Border, Color, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.pagebreak import Break

from certpdf.config import CERTIFICATES, PASS_FAIL, default_sheet_names
from certpdf.engines import LibreOfficeEngine
from certpdf.excel_output import build_workbooks
from certpdf.sheet_copy import tab_title
from certpdf.workbook_info import auto_height_rows

TABS = ["Device data", "ECG.certificate", "NIBP.certificate", "ECG.pass-fail test sheet", "NIBP.pass-fail sheet"]


class Recorder:
    def __init__(self):
        self.logs, self.statuses = [], {}

    def log(self, level, text):
        self.logs.append((level, text))

    def progress(self, done, total, text):
        assert 0 <= done <= total

    def file_status(self, index, state, text):
        self.statuses[index] = state


def make_template(path: Path, device: str, tabs=TABS) -> Path:
    """A workbook shaped like the real templates (Arial 10 default font, merged cells, print setup...)."""
    wb = Workbook()
    arial = Font(name="Arial", sz=10)
    wb._fonts = IndexedList([arial])
    wb._named_styles["Normal"].font = arial
    wb.remove(wb.active)
    thin = Side(style="thin")
    for tab in tabs:
        ws = wb.create_sheet(tab)
        ws["A1"] = f"{device} {tab}"
        ws["A1"].font = Font(name="Arial Narrow", sz=14, bold=True)
        ws.merge_cells("A1:D1")
        ws["A1"].border = Border(left=thin, top=thin, bottom=thin)
        ws["D1"].border = Border(right=thin, top=thin, bottom=thin)
        for row in range(3, 40):
            ws.cell(row, 1, f"Check {row}")
            ws.cell(row, 2, "PASS")
        ws.column_dimensions["A"].width = 13.7
        ws.column_dimensions["B"].width = 4.4
        ws.row_dimensions[3].height = 30
        ws.print_area = "A1:D39"
        ws.page_setup.scale = 82
        ws.oddFooter.center.text = "Code No MECL-TR-05"
        ws.row_breaks.append(Break(id=20))
        # "Hide these cells" rule like the real pass/fail sheets: number format ;;; and no fill,
        # which is how openpyxl reads Excel's patternType="none".
        hide = DifferentialStyle(numFmt=NumberFormat(numFmtId=167, formatCode=";;;"),
                                 fill=PatternFill(bgColor=Color(auto=True)))
        ws.conditional_formatting.add("B30:C35", Rule(type="expression", formula=['$A$1="x"'], dxf=hide))
        ws.conditional_formatting.add("B36:C37", FormulaRule(formula=['$A$1="y"'],
                                                             fill=PatternFill("solid", bgColor="FFFF0000")))
        ws.conditional_formatting.add("B38:C39", FormulaRule(formula=['$A$1="z"'], font=Font(color="FFFF0000")))
    wb.save(path)
    return path


@pytest.fixture
def template(tmp_path):
    def factory(name, device=None, **kwargs):
        return make_template(tmp_path / name, device or Path(name).stem, **kwargs)
    return factory


def build(files, tmp_path, **kwargs):
    outputs = {CERTIFICATES: tmp_path / "out" / "Certificates.xlsx",
               PASS_FAIL: tmp_path / "out" / "Pass-Fail Test Sheets.xlsx"}
    recorder = Recorder()
    result = build_workbooks(files, default_sheet_names(), outputs, recorder, **kwargs)
    return result, outputs, recorder


def test_tabs_in_order_one_per_sheet(template, tmp_path):
    files = [template("F23-AGH020-0726.xlsx"), template("F23-AGH019-0626.xlsx")]
    result, outputs, recorder = build(files, tmp_path)

    certs = load_workbook(outputs[CERTIFICATES])
    assert certs.sheetnames == ["F23-AGH020-0726 ECG cert", "F23-AGH020-0726 NIBP cert",
                                "F23-AGH019-0626 ECG cert", "F23-AGH019-0626 NIBP cert"]
    assert certs["F23-AGH019-0626 NIBP cert"]["A1"].value == "F23-AGH019-0626 NIBP.certificate"
    passfail = load_workbook(outputs[PASS_FAIL])
    assert passfail.sheetnames == ["F23-AGH020-0726 ECG pass-fail", "F23-AGH020-0726 NIBP pass-fail",
                                   "F23-AGH019-0626 ECG pass-fail", "F23-AGH019-0626 NIBP pass-fail"]
    assert result.outputs[CERTIFICATES].tabs == 4
    assert recorder.statuses == {0: "done", 1: "done"}


def test_layout_and_print_setup_are_kept(template, tmp_path):
    _, outputs, _ = build([template("dev.xlsx")], tmp_path)
    ws = load_workbook(outputs[PASS_FAIL])["dev ECG pass-fail"]

    assert [str(r) for r in ws.merged_cells.ranges] == ["A1:D1"]
    assert ws["A1"].font.name == "Arial Narrow" and ws["A1"].font.b
    assert ws["A1"].border.left.style == "thin" and ws["D1"].border.right.style == "thin"
    assert ws.column_dimensions["A"].width == pytest.approx(13.7)
    assert ws.column_dimensions["B"].width == pytest.approx(4.4)
    assert ws.row_dimensions[3].height == 30
    assert ws.print_area == "'dev ECG pass-fail'!$A$1:$D$39"
    assert ws.page_setup.scale == 82
    assert ws.page_setup.paperSize == 9  # A4 when the tab didn't say
    assert ws.oddFooter.center.text == "Code No MECL-TR-05"
    assert [b.id for b in ws.row_breaks.brk] == [20]
    rules = [rule for cf in ws.conditional_formatting for rule in cf.rules]
    assert len(rules) == 3


def test_no_fill_conditional_format_does_not_become_black(template, tmp_path):
    _, outputs, _ = build([template("dev.xlsx")], tmp_path)
    with zipfile.ZipFile(outputs[PASS_FAIL]) as zf:
        styles = zf.read("xl/styles.xml").decode()
    dxfs = re.findall(r"<dxf>.*?</dxf>", styles, re.S)
    hide = [d for d in dxfs if ";;;" in d]
    # The "no fill" rule keeps hiding the values but must not become a solid (black) fill...
    assert hide and "<fill>" not in hide[0]
    # ...while a real fill colour is kept.
    assert any("<fill>" in d and "FFFF0000" in d for d in dxfs)


def test_default_font_is_kept_so_column_widths_match(template, tmp_path):
    _, outputs, _ = build([template("dev.xlsx")], tmp_path)
    with zipfile.ZipFile(outputs[CERTIFICATES]) as zf:
        styles = zf.read("xl/styles.xml").decode()
    first_font = re.search(r"<fonts[^>]*>\s*<font>(.*?)</font>", styles, re.S).group(1)
    assert 'val="Arial"' in first_font and 'val="10"' in first_font


def test_missing_tabs_and_unreadable_files(template, tmp_path):
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"not a workbook")
    files = [template("full.xlsx"), template("no_nibp.xlsx", tabs=["ECG Certificate", "ECG Pass Fail"]), broken]
    result, outputs, recorder = build(files, tmp_path)

    assert len(load_workbook(outputs[CERTIFICATES]).sheetnames) == 3
    assert result.workbooks[1].missing == ["NIBP certificate", "NIBP pass/fail"]
    assert result.workbooks[2].problems and "could not be read" in result.workbooks[2].problems[0]
    assert recorder.statuses == {0: "done", 1: "warning", 2: "error"}


def test_group_without_tabs_is_not_saved(template, tmp_path):
    result, outputs, recorder = build([template("dev.xlsx", tabs=["ECG.certificate"])], tmp_path)
    assert result.outputs[PASS_FAIL] is None and not outputs[PASS_FAIL].exists()
    assert any("no Excel file saved" in text for level, text in recorder.logs if level == "warning")


def test_tab_titles_are_valid_and_unique():
    used = set()
    assert tab_title("F23-AGH019-0626", "ECG cert", used) == "F23-AGH019-0626 ECG cert"
    long = tab_title("A very long device file name that goes on", "NIBP pass-fail", used)
    assert len(long) <= 31 and long.endswith("NIBP pass-fail")
    assert tab_title("x[1]:y", "ECG cert", used) == "x-1--y ECG cert"
    assert tab_title("F23-AGH019-0626", "ECG cert", used) == "F23-AGH019-0626 ECG cert (2)"


def test_auto_height_rows(template, tmp_path):
    path = template("dev.xlsx")
    # openpyxl always writes customHeight; make row 3 look like a height Excel calculated itself.
    patched = tmp_path / "patched.xlsx"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(patched, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet2.xml":
                data = re.sub(rb'(<row r="3"[^>]*?) customHeight="1"', rb"\1", data)
            zout.writestr(item, data)
    rows = auto_height_rows(patched, ["ECG.certificate", "NIBP.certificate"])
    assert rows == {"ECG.certificate": {3}, "NIBP.certificate": set()}


soffice = LibreOfficeEngine.find()


@pytest.mark.skipif(soffice is None, reason="LibreOffice is not installed")
def test_values_not_formulas(tmp_path):
    # LibreOffice saves the calculated values, like Excel does.
    wb = Workbook()
    wb.active.title = "Device data"
    wb["Device data"]["A1"] = "F23-AGH019-0626"
    cert = wb.create_sheet("ECG.certificate")
    cert["A1"] = "='Device data'!A1"
    cert["A2"] = "=6*7"
    raw = tmp_path / "raw.xlsx"
    wb.save(raw)
    calc_dir = tmp_path / "calc"
    subprocess.run([soffice, f"-env:UserInstallation={(tmp_path / 'profile').as_uri()}", "--headless",
                    "--convert-to", "xlsx", "--outdir", str(calc_dir), str(raw)],
                   check=True, capture_output=True, timeout=180)
    _, outputs, _ = build([calc_dir / "raw.xlsx"], tmp_path)

    ws = load_workbook(outputs[CERTIFICATES]).worksheets[0]
    assert ws["A1"].value == "F23-AGH019-0626"
    assert ws["A2"].value == 42
